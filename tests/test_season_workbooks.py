"""Integration coverage for isolated, chronological season workbooks."""

from datetime import date

import pytest
from openpyxl import load_workbook

from softball_statistics.cli import CLI
from softball_statistics.exporters.excel_exporter import ExcelExporter, ExcelExportError
from softball_statistics.models import Game, League, PlateAppearance, Player, Team, Week
from softball_statistics.parsers.csv_parser import CSVParser
from softball_statistics.repository.sqlite import SQLiteRepository
from softball_statistics.use_cases import CalculateStatsUseCase


def add_game(repo, league, season, day, number, outcomes, team="Cyclones"):
    """Persist real games in deliberately nonchronological insertion order."""
    league_id = repo.save_league(League(None, league, season))
    team_id = repo.save_team(Team(None, league_id, team))
    player_id = repo.save_player(Player(None, team_id, "Alex"))
    game_date = date.fromisoformat(day)
    week_id = repo.save_week(Week(None, league_id, number, game_date, game_date))
    game_id = repo.save_game(Game(None, week_id, team_id, game_date, number))
    for outcome, bases in outcomes:
        repo.save_plate_appearance(
            PlateAppearance(None, player_id, game_id, outcome, bases=bases)
        )


@pytest.fixture
def season_repo(tmp_path):
    repo = SQLiteRepository(str(tmp_path / "stats.db"))
    # Same player/team names in another league must never leak into Fray totals.
    add_game(repo, "Scottsdale", "Winter 2026", "2026-02-01", 1, [("3B", 3)])
    add_game(repo, "Fray", "Summer 2026", "2026-06-01", 1, [("HR", 4)])
    add_game(repo, "Fray", "Spring 2026", "2026-04-01", 2, [("2B", 2)])
    add_game(repo, "Fray", "Winter 2026", "2026-02-01", 1, [("1B", 1), ("O", 0)])
    add_game(repo, "Fray", "Spring 2026", "2026-04-01", 1, [("BB", 0), ("O", 0)])
    return repo


def sheet_records(workbook, name):
    rows = workbook[name].iter_rows(values_only=True)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows if row[0] is not None]


def export(repo, tmp_path):
    return ExcelExporter(repo).export(
        {}, str(tmp_path / "stats.xlsx"), use_case=CalculateStatsUseCase(repo)
    )


def test_separate_workbooks_and_historical_totals(season_repo, tmp_path):
    paths = export(season_repo, tmp_path)
    expected = {
        "stats-fray-winter_2026.xlsx",
        "stats-fray-spring_2026.xlsx",
        "stats-fray-summer_2026.xlsx",
        "stats-scottsdale-winter_2026.xlsx",
    }
    assert {p.name for p in tmp_path.glob("*.xlsx")} == expected
    assert {str(tmp_path / name) for name in expected} == set(paths)

    with_workbooks = [
        ("winter", 2, 1, "0.500", 1),
        ("spring", 5, 2, "0.500", 3),
        ("summer", 6, 3, "0.600", 4),
    ]
    for season, pa, hits, ba, games in with_workbooks:
        workbook = load_workbook(tmp_path / f"stats-fray-{season}_2026.xlsx")
        try:
            summary = sheet_records(workbook, "League Summary")
            summary_sheet = workbook["League Summary"]
            assert len(summary) == 1
            assert summary_sheet["I1"].border.bottom.style == "thin"
            assert summary_sheet["J1"].border.bottom.style == "thin"
            assert summary[0]["League"] == "Fray"
            assert summary[0]["Games Played"] == games
            assert summary[0]["Team BA"] == ba
            player = sheet_records(workbook, "Player Summary")[0]
            assert (player["Player"], player["PA"], player["H"], player["BA"]) == (
                "Alex",
                pa,
                hits,
                ba,
            )
            assert player["3B"] == 0
            assert player["HR"] == (1 if season == "summer" else 0)
            game_tabs = [name for name in workbook.sheetnames if " Game " in name]
            assert len(game_tabs) == (2 if season == "spring" else 1)
            assert all(season.title() in name for name in game_tabs)
            if season == "spring":
                assert game_tabs == [
                    "Cyclo Spring 2026 Game 2",
                    "Cyclo Spring 2026 Game 1",
                ]
                totals = sheet_records(workbook, "Cyclo Spring 2026 Total")[0]
                assert (totals["PA"], totals["H"]) == (3, 1)
                assert player["OBP"] == "0.600"
                assert player["SLG"] == "0.750"
                assert player["OPS"] == "1.350"
        finally:
            workbook.close()


def test_backfill_refreshes_later_totals_but_does_not_leak_future_games(
    season_repo, tmp_path
):
    export(season_repo, tmp_path)
    # A late-entered game from the earlier season affects all later snapshots.
    add_game(season_repo, "Fray", "Winter 2026", "2026-02-08", 2, [("1B", 1)])
    # Overlapping season dates require filtering individual games, not whole seasons.
    add_game(season_repo, "Fray", "Winter 2026", "2026-07-01", 3, [("HR", 4)])
    paths = export(season_repo, tmp_path)
    assert len(paths) == 4
    workbook = load_workbook(tmp_path / "stats-fray-spring_2026.xlsx")
    try:
        player = sheet_records(workbook, "Player Summary")[0]
        assert (player["PA"], player["H"], player["HR"]) == (6, 3, 0)
        assert sheet_records(workbook, "League Summary")[0]["Games Played"] == 4
        assert len([s for s in workbook.sheetnames if " Game " in s]) == 2
    finally:
        workbook.close()


def test_cli_reports_actual_workbooks_for_all_teams(season_repo, tmp_path, capsys):
    source = tmp_path / "fray-storm-sp-03_2026-04-08.csv"
    source.write_text("Player Name,Attempt\nSam,1B\n")
    CLI(season_repo, season_repo, CSVParser(), ExcelExporter(season_repo)).run(
        ["--file", str(source), "--output", str(tmp_path / "stats.xlsx")]
    )
    output = capsys.readouterr().out
    assert "stats-fray-spring_2026.xlsx" in output
    workbook = load_workbook(tmp_path / "stats-fray-spring_2026.xlsx")
    try:
        assert "Cyclo Spring 2026 Game 1" in workbook.sheetnames
        assert "Storm Spring 2026 Game 3" in workbook.sheetnames
        assert {r["Team"] for r in sheet_records(workbook, "League Summary")} == {
            "Cyclones",
            "Storm",
        }
    finally:
        workbook.close()


def test_colliding_team_abbreviations_and_long_seasons_keep_every_sheet(tmp_path):
    repo = SQLiteRepository(str(tmp_path / "stats.db"))
    season = "An Exceptionally Long Summer Season 2026"
    add_game(repo, "Fray", season, "2026-06-01", 1, [("1B", 1)], team="Blue Bears")
    add_game(repo, "Fray", season, "2026-06-01", 1, [("2B", 2)], team="Brown Bears")
    paths = export(repo, tmp_path)
    assert len(paths) == 1
    workbook = load_workbook(paths[0])
    try:
        assert len(workbook.sheetnames) == 9
        assert len({name.casefold() for name in workbook.sheetnames}) == 9
        assert all(len(name) <= 31 for name in workbook.sheetnames)
        # The last two sheets are separate game tabs, not overwritten by initials.
        players = [
            sheet_records(workbook, name)[0] for name in workbook.sheetnames[-2:]
        ]
        assert [(p["1B"], p["2B"]) for p in players] == [(0, 1), (1, 0)]
    finally:
        workbook.close()


def test_sanitized_filename_collisions_keep_distinct_leagues(tmp_path):
    repo = SQLiteRepository(str(tmp_path / "stats.db"))
    add_game(repo, "A/B", "Summer 2026", "2026-06-01", 1, [("1B", 1)])
    add_game(repo, "A B", "Summer 2026", "2026-06-01", 1, [("2B", 2)])
    paths = export(repo, tmp_path)
    assert len(set(paths)) == 2
    assert all("stats-a_b-summer_2026-" in path for path in paths)


def test_no_workbook_for_empty_season(tmp_path):
    repo = SQLiteRepository(str(tmp_path / "stats.db"))
    repo.save_league(League(None, "Fray", "Winter 2026"))
    assert export(repo, tmp_path) == []
    assert not list(tmp_path.glob("*.xlsx"))


def test_export_reports_filesystem_errors(season_repo, tmp_path):
    blocked = tmp_path / "blocked"
    blocked.write_text("Not a directory")
    with pytest.raises(ExcelExportError, match="Failed to export season workbooks"):
        ExcelExporter(season_repo).export({}, str(blocked / "stats.xlsx"))
